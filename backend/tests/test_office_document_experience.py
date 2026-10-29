"""Office ingestion regression: Word, Excel, slides and upload-to-index path."""
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pptx import Presentation

from app.parsers import parse_document


def make_word_file() -> bytes:
    document = Document()
    document.add_heading("Quarterly review", level=1)
    document.add_paragraph("This report tracks client work.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Category"
    table.cell(0, 1).text = "Revenue"
    document.sections[0].header.paragraphs[0].text = "Internal reporting"
    document.sections[0].footer.paragraphs[0].text = "Confidential footer"
    data = BytesIO()
    document.save(data)
    return data.getvalue()


def make_excel_file() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Budget 2026"
    sheet["A1"] = "Category"
    sheet["B1"] = "Amount"
    sheet["A2"] = "Software"
    sheet["B2"] = 1200
    sheet["A3"] = "Total"
    sheet["B3"] = "=SUM(B2)"  # OpenPyXL cannot evaluate formulas.
    sheet["D3"] = "日本語"
    other = workbook.create_sheet("Another sheet")
    other["A1"] = "Operating cost"
    data = BytesIO()
    workbook.save(data)
    return data.getvalue()


def make_powerpoint_file() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Launch update"
    slide.placeholders[1].text = "The rollout is scheduled for November."
    slide.notes_slide.notes_text_frame.text = "Verify launch milestones."
    data = BytesIO()
    presentation.save(data)
    return data.getvalue()


def test_office_parser_extracts_word_tables_header_footer(tmp_path) -> None:
    path = tmp_path / "report.docx"
    path.write_bytes(make_word_file())
    result = parse_document(path)
    content = "\n".join(unit.content for unit in result.units)
    assert "Quarterly review" in content
    assert "Category" in content and "Revenue" in content
    assert "Internal reporting" in content
    assert "Confidential footer" in content


def test_excel_parser_preserves_sheet_cell_references_and_formulas(tmp_path) -> None:
    path = tmp_path / "budget.xlsx"
    path.write_bytes(make_excel_file())
    result = parse_document(path)
    assert len(result.units) == 2
    sheet = result.units[0]
    assert sheet.source_label == "Budget 2026"
    assert "A2: Software" in sheet.content
    assert "B2: 1200" in sheet.content
    assert "B3: =SUM(B2)" in sheet.content
    assert "D3: 日本語" in sheet.content
    assert sheet.metadata["extracted_rows"] == 3
    assert sheet.metadata["truncated"] is False
    assert "Operating cost" in result.units[1].content


def test_pptx_parser_reads_slide_and_speaker_notes(tmp_path) -> None:
    path = tmp_path / "presentation.pptx"
    path.write_bytes(make_powerpoint_file())
    result = parse_document(path)
    assert result.page_count == 1
    assert "Launch update" in result.units[0].content
    assert "The rollout is scheduled for November" in result.units[0].content
    assert "Verify launch milestones" in result.units[0].content
    assert result.units[0].source_label == "Slide 1"


def _access(client: TestClient) -> dict[str, str]:
    created = client.post("/api/v1/auth/register", json={
        "email": "office-reader@example.com",
        "username": "office_reader",
        "password": "OfficePassword123!",
    })
    assert created.status_code == 201
    login = client.post("/api/v1/auth/login", json={
        "email": "office-reader@example.com",
        "password": "OfficePassword123!",
    })
    assert login.status_code == 200
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


def test_real_office_upload_process_and_authenticated_units(client: TestClient) -> None:
    headers = _access(client)
    fixtures = (
        ("finance.xlsx", make_excel_file()),
        ("document.docx", make_word_file()),
        ("deck.pptx", make_powerpoint_file()),
    )
    for filename, payload in fixtures:
        response = client.post("/api/v1/documents/upload", headers=headers, files={
            "file": (filename, payload, "application/octet-stream"),
        })
        assert response.status_code == 201, response.text
        document_id = response.json()["id"]
        processed = client.post(
            f"/api/v1/documents/{document_id}/process", headers=headers,
        )
        assert processed.status_code == 200, processed.text
        assert processed.json()["status"] == "ready"
        units = client.get(
            f"/api/v1/documents/{document_id}/units?limit=30&offset=0",
            headers=headers,
        )
        assert units.status_code == 200
        assert units.json()["total"] > 0
        assert any(row["content"] for row in units.json()["items"])
        anon = client.get(f"/api/v1/documents/{document_id}/units")
        assert anon.status_code == 401


def test_excel_parser_caps_extreme_extracted_content(tmp_path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    for row in range(1, 140):
        sheet.cell(row, 1, "x" * 1200)
    path = tmp_path / "large.xlsx"
    workbook.save(path)
    result = parse_document(path)
    assert result.units[0].metadata["truncated"] is True
    assert len(result.units[0].content) <= 120_000
