"""End-to-end artifacts authorization, real bytes, and formula safety."""
from io import BytesIO
from types import SimpleNamespace

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation

from app.services.artifact_rendering import render_artifact


def _auth(client: TestClient, username: str) -> dict[str, str]:
    email = username + "@example.com"
    assert client.post("/api/v1/auth/register", json={
        "username": username, "email": email, "password": "SecurePass123!",
    }).status_code == 201
    token_response = client.post("/api/v1/auth/login", json={
        "email": email, "password": "SecurePass123!",
    })
    assert token_response.status_code == 200
    return {"Authorization": "Bearer " + token_response.json()["access_token"]}


def test_all_formats_render_real_files():
    markdown = (
        "# RAG Research\n\n## Results\n"
        "| Year | Cost |\n| --- | --- |\n| 2026 | 350 |\n"
        "| 2027 | =HYPERLINK(1) |\n"
        "## Conclusion\n- Citation validated\n- Private workspace\n"
    )
    assert render_artifact(markdown, "md").startswith(b"# RAG")
    docx = Document(BytesIO(render_artifact(markdown, "docx")))
    assert any("RAG Research" in p.text for p in docx.paragraphs)
    xlsx = load_workbook(BytesIO(render_artifact(markdown, "xlsx")), data_only=False)
    assert xlsx.active["B5"].data_type == "s"
    assert xlsx.active["B5"].value.startswith("'=HYPERLINK")
    slides = Presentation(BytesIO(render_artifact(markdown, "pptx")))
    assert len(slides.slides) >= 2
    assert render_artifact(markdown, "pdf").startswith(b"%PDF")
    for kind in ["xlsx", "pptx", "docx"]:
        assert render_artifact(markdown, kind).startswith(b"PK")


def test_artifacts_auth_owner_scope_versions_and_download(client: TestClient):
    url = "/api/v1/artifacts"
    assert client.get(url).status_code == 401
    assert client.post(url + "/render", json={
        "title": "Report", "format": "md", "content": "# Secure",
    }).status_code == 401
    alpha = _auth(client, "artifactalpha")
    beta = _auth(client, "artifactbeta")
    response = client.post(url + "/render", headers=alpha, json={
        "title": "Report", "format": "md", "content": "# Secure report",
    })
    assert response.status_code == 201
    item = response.json()
    assert item["filename"].endswith(".md")
    assert item["content"] == "# Secure report"
    artifact_id = item["id"]
    assert client.get(url, headers=alpha).json()["total"] == 1
    assert client.get(url, headers=beta).json()["total"] == 0
    assert client.get(url + "/" + artifact_id, headers=beta).status_code == 404
    assert client.get(url + "/" + artifact_id + "/download", headers=beta).status_code == 404
    assert client.delete(url + "/" + artifact_id, headers=beta).status_code == 404
    download = client.get(url + "/" + artifact_id + "/download", headers=alpha)
    assert download.status_code == 200
    assert download.content == b"# Secure report"
    assert download.headers["cache-control"] == "private, no-store"
    assert "attachment" in download.headers["content-disposition"]
    assert client.delete(url + "/" + artifact_id, headers=alpha).status_code == 204
    assert client.get(url + "/" + artifact_id, headers=alpha).status_code == 404


def test_deterministic_local_mode_fails_closed(client: TestClient):
    alpha = _auth(client, "artifactlocal")
    response = client.post("/api/v1/artifacts/generate", headers=alpha, json={
        "title": "Draft", "format": "pptx",
        "prompt": "Write a professional presentation about RAG.",
    })
    assert response.status_code == 503
    assert client.get("/api/v1/artifacts", headers=alpha).json()["total"] == 0


def test_real_provider_output_generates_docx(client: TestClient, monkeypatch):
    alpha = _auth(client, "artifactwriter")
    class Provider:
        def generate(self, **kwargs):
            assert "Write a research report" in kwargs["input_text"]
            return SimpleNamespace(
                text="# Research\n\n## Findings\n- Real content generated",
                provider_name="test-provider",
            )
    monkeypatch.setattr(
        "app.api.artifacts.settings.LLM_PROVIDER", "openai",
    )
    monkeypatch.setattr(
        "app.api.artifacts.create_configured_llm_provider",
        lambda: Provider(),
    )
    response = client.post("/api/v1/artifacts/generate", headers=alpha, json={
        "title": "Research", "format": "docx",
        "prompt": "Write a research report about systems and AI.",
    })
    assert response.status_code == 201
    item = response.json()
    assert item["provider_name"] == "test-provider"
    data = client.get(
        "/api/v1/artifacts/" + item["id"] + "/download", headers=alpha,
    )
    assert data.status_code == 200
    doc = Document(BytesIO(data.content))
    assert any("Real content generated" in p.text for p in doc.paragraphs)


def test_no_external_users_document_leak(client: TestClient):
    url = "/api/v1/artifacts/generate"
    alpha = _auth(client, "artifactowner")
    beta = _auth(client, "artifactother")
    uploaded = client.post("/api/v1/documents/upload", headers=alpha, files={
        "file": ("secret.md", b"# Private source document", "text/markdown"),
    })
    assert uploaded.status_code == 201
    document_id = uploaded.json()["id"]
    result = client.post(url, headers=beta, json={
        "title": "Attempt", "format": "docx",
        "prompt": "Generate report from another user document.",
        "document_ids": [document_id],
    })
    assert result.status_code == 422
    assert client.get("/api/v1/artifacts", headers=beta).json()["total"] == 0


def test_bad_identifiers_not_usable(client: TestClient):
    alpha = _auth(client, "artifactpath")
    assert client.get(
        "/api/v1/artifacts/../../etc/passwd", headers=alpha,
        follow_redirects=False,
    ).status_code in (404, 405)
    assert client.get(
        "/api/v1/artifacts/not-a-uuid", headers=alpha,
    ).status_code == 404


def test_grounded_artifact_reuses_scoped_rag_and_citations(
    client: TestClient, db_session, monkeypatch,
):
    from sqlalchemy import select
    from app.models.document import Document

    owner = _auth(client, "artifactgrounded")
    uploaded = client.post("/api/v1/documents/upload", headers=owner, files={
        "file": ("my-source.md", b"# Private evidence\nRAG content.", "text/markdown"),
    })
    assert uploaded.status_code == 201
    doc_id = uploaded.json()["id"]
    document = db_session.scalar(select(Document).where(Document.id == doc_id))
    document.status = "ready"
    db_session.commit()
    calls = []

    def fake_answer_question(**kwargs):
        calls.append(kwargs)
        return SimpleNamespace(
            is_refusal=False,
            citations=(SimpleNamespace(source_id="S1"),),
            answer_text="## Verified evidence\nA fact supported by [S1].",
            provider_name="test-rag",
            citation_count=1,
        )

    monkeypatch.setattr("app.api.artifacts.answer_question", fake_answer_question)
    response = client.post("/api/v1/artifacts/generate", headers=owner, json={
        "title": "Grounded", "format": "md",
        "prompt": "Summarize the private evidence from my documents.",
        "document_ids": [doc_id],
    })
    assert response.status_code == 201
    data = response.json()
    assert data["citation_count"] == 1
    assert data["source_document_ids"] == [doc_id]
    assert "[S1]" in data["content"]
    assert len(calls) == 1
    assert calls[0]["document_ids"] == (doc_id,)
    assert calls[0]["user_id"] == document.user_id
