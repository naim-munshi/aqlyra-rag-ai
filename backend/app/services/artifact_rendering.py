"""Safe, bounded in-process renderers; no macros, shell or formula evaluation."""
from __future__ import annotations

import io
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches as DocxInches, Pt as DocxPt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

MIMES = {
    "md": "text/markdown; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "pdf": "application/pdf",
}
MAX_CONTENT_CHARS = 24_000
MAX_TABLE_ROWS = 450
MAX_TABLE_COLUMNS = 12
MAX_SLIDES = 25


def _lines(content: str) -> list[str]:
    lines = content.splitlines()
    if len(lines) > 450:
        raise ValueError("Artifact exceeds the maximum 450 source lines")
    return [line.strip() for line in lines]


def _strip(text: str) -> str:
    return re.sub(r"[*_]", "", text).replace(chr(96), "").strip()


def _table_cells(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    if len(cells) > MAX_TABLE_COLUMNS:
        raise ValueError("Artifact table exceeds 12 columns")
    return [_strip(value) for value in cells]


def _is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\|?[\s|:\-]+\|?", line)) and "-" in line


def _is_formula(text: str) -> bool:
    # Avoid spreadsheet formula evaluation on untrusted generated content.
    return bool(text.lstrip()) and text.lstrip()[0] in "=+-@"


def render_markdown(content: str) -> bytes:
    return content.encode("utf-8")


def render_docx(content: str) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = DocxInches(0.7)
    section.bottom_margin = DocxInches(0.7)
    style = doc.styles["Normal"]
    style.font.size = DocxPt(10)
    for line in _lines(content):
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(_strip(line[4:]), level=3)
        elif line.startswith("## "):
            doc.add_heading(_strip(line[3:]), level=2)
        elif line.startswith("# "):
            doc.add_heading(_strip(line[2:]), level=1)
        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(_strip(line[2:]), style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", line):
            doc.add_paragraph(_strip(re.sub(r"^\d+[.)]\s+", "", line)), style="List Number")
        elif line.startswith("|"):
            doc.add_paragraph(" | ".join(_table_cells(line)))
        else:
            doc.add_paragraph(_strip(line))
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def render_xlsx(content: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Aqlyra Report"
    row = 1
    for line in _lines(content):
        if not line or _is_separator(line):
            continue
        if row > MAX_TABLE_ROWS:
            raise ValueError("Artifact spreadsheet exceeds the row limit")
        items = _table_cells(line) if line.startswith("|") else [_strip(line)]
        for col, value in enumerate(items[:MAX_TABLE_COLUMNS], 1):
            if _is_formula(value):
                value = "'" + value
            cell = sheet.cell(row=row, column=col, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if line.startswith("# ") or (line.startswith("|") and row == 1):
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="203C70")
        row += 1
    sheet.freeze_panes = "A2"
    for col in "ABCDEFGHIJKL":
        sheet.column_dimensions[col].width = 27
    workbook.calculation.fullCalcOnLoad = False
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_pptx(content: str) -> bytes:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    lines = _lines(content)
    title = next((line[2:].strip() for line in lines if line.startswith("# ")), "Aqlyra Presentation")
    intro = presentation.slides.add_slide(presentation.slide_layouts[0])
    intro.shapes.title.text = _strip(title)
    intro.placeholders[1].text = "Generated with Aqlyra"
    current_title = "Overview"
    bullets: list[str] = []

    def flush() -> None:
        nonlocal bullets
        if not bullets:
            return
        if len(presentation.slides) >= MAX_SLIDES:
            raise ValueError("Presentation exceeds 25 slides")
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        if len(current_title) > 100:
            raise ValueError("Slide heading is too long")
        slide.shapes.title.text = _strip(current_title)
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, item in enumerate(bullets[:7]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            if len(item) > 200:
                raise ValueError("Slide bullet is too long for readable layout")
            paragraph.text = _strip(item)
            paragraph.level = 0
            paragraph.font.size = Pt(20)
        bullets = []

    for line in lines:
        if line.startswith("## "):
            flush()
            current_title = line[3:]
        elif line.startswith("# "):
            continue
        elif line:
            if _is_separator(line):
                continue
            bullets.append(re.sub(r"^(?:[-*]|\d+[.)])\s+", "", line))
            if len(bullets) >= 7:
                flush()
    flush()
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _pdf_font() -> str:
    for path in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ):
        try:
            pdfmetrics.registerFont(TTFont("AqlyraBody", path))
            return "AqlyraBody"
        except (OSError, ValueError):
            continue
    return "Helvetica"


def render_pdf(content: str) -> bytes:
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=A4, topMargin=48, bottomMargin=48,
        leftMargin=48, rightMargin=48,
    )
    font_name = _pdf_font()
    styles = getSampleStyleSheet()
    regular = ParagraphStyle(
        "AqlyraRegular", parent=styles["Normal"],
        fontName=font_name, fontSize=10, leading=15,
        alignment=TA_LEFT, spaceAfter=8,
    )
    heading = ParagraphStyle(
        "AqlyraHeading", parent=regular,
        fontSize=16, leading=21, textColor=colors.HexColor("#203C70"),
        spaceBefore=12,
    )
    story = []
    for line in _lines(content):
        if not line:
            story.append(Spacer(1, 5))
            continue
        selected = heading if line.startswith("#") else regular
        text = re.sub(r"^#{1,3}\s+", "", line)
        story.append(Paragraph(escape(_strip(text)), selected))
    if not story:
        story.append(Paragraph("Aqlyra Document", regular))
    doc.build(story)
    return output.getvalue()


RENDERERS = {
    "md": render_markdown,
    "docx": render_docx,
    "xlsx": render_xlsx,
    "pptx": render_pptx,
    "pdf": render_pdf,
}


def render_artifact(content: str, kind: str) -> bytes:
    if kind not in RENDERERS:
        raise ValueError("Unsupported artifact format")
    if not content.strip() or len(content) > MAX_CONTENT_CHARS:
        raise ValueError("Artifact content is empty or too large")
    data = RENDERERS[kind](content)
    if not data or len(data) > 5 * 1024 * 1024:
        raise ValueError("Rendered file exceeds the 5 MiB size limit")
    return data
